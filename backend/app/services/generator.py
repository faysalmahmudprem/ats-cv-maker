"""DOCX renderer for the shared CV layout.

This is a PRESENTATION-only module: every content decision (what appears,
in what order, which heading text, which parts join into a headline) is
made by app/services/layout.py so the DOCX and PDF outputs can never
drift apart. The formatting code (fonts, accent color, heading borders,
bullets, clickable OOXML hyperlinks) is carried over unchanged from the
original generator so existing output quality is preserved.

Public API:
    build_cv(cv, template_key, profile_key) -> io.BytesIO
    generate_docx_bytes(cv, template_key, profile_key) -> bytes
"""

from __future__ import annotations

import io
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.profiles import Profile
from app.services.layout import Block, Layout, build_layout
from app.templates import TemplateStyle, get_template

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _accent(style: TemplateStyle) -> RGBColor:
    return RGBColor.from_string(style.accent_hex)


def _align(header_align: str):
    return (
        WD_ALIGN_PARAGRAPH.CENTER
        if header_align == "center"
        else WD_ALIGN_PARAGRAPH.LEFT
    )


def add_hyperlink(
    paragraph: Any,
    text: str,
    url: str,
    style: TemplateStyle,
) -> None:
    """Clickable link, readable text. Renders plain text if url empty."""
    if not url:
        r = paragraph.add_run(text)
        r.font.size = Pt(style.contact_pt)
        r.font.name = style.font
        return
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), style.accent_hex)
    props.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    props.append(u)
    run.append(props)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def style_document(doc: Document, style: TemplateStyle) -> None:
    top, bottom, left, right = style.margins_in
    s = doc.sections[0]
    s.top_margin = Inches(top)
    s.bottom_margin = Inches(bottom)
    s.left_margin = Inches(left)
    s.right_margin = Inches(right)
    n = doc.styles["Normal"]
    n.font.name = style.font
    n.font.size = Pt(style.body_pt)
    n.paragraph_format.space_before = Pt(0)
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    b = doc.styles["List Bullet"]
    b.font.name = style.font
    b.font.size = Pt(style.body_pt)


def heading(doc: Document, text: str, style: TemplateStyle) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(style.heading_pt)
    r.font.color.rgb = _accent(style)
    r.font.name = style.font
    p.paragraph_format.space_before = Pt(style.heading_space_before)
    p.paragraph_format.space_after = Pt(style.heading_space_after)
    if style.heading_border:
        pr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "6")
        bot.set(qn("w:color"), style.accent_hex)
        bd.append(bot)
        pr.append(bd)


def bullet(doc: Document, text: str, style: TemplateStyle) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    r.font.name = style.font
    r.font.size = Pt(style.body_pt)


# ---------------------------------------------------------------------------
# Block renderers — presentation only; content decided by layout.py.
# ---------------------------------------------------------------------------


def _render_para(doc: Document, block: Block, style: TemplateStyle) -> None:
    p = doc.add_paragraph()
    r = p.add_run(block.text)
    r.font.name = style.font
    r.font.size = Pt(style.body_pt)


def _render_skillline(doc: Document, block: Block, style: TemplateStyle) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    lb = p.add_run((block.label + ": ") if block.label else "")
    lb.bold = True
    lb.font.name = style.font
    lb.font.size = Pt(style.body_pt)
    bd = p.add_run(block.text)
    bd.font.name = style.font
    bd.font.size = Pt(style.body_pt)


def _render_job_or_edu(
    doc: Document, block: Block, style: TemplateStyle, kind: str
) -> None:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(
        style.entry_space_before if kind == "job" else style.edu_space_before
    )
    h.paragraph_format.space_after = Pt(0)
    t = h.add_run(block.headline)
    t.bold = True
    t.font.size = Pt(style.entry_pt)
    t.font.name = style.font
    if block.meta:
        m = doc.add_paragraph()
        if kind == "job":
            m.paragraph_format.space_after = Pt(2)
        x = m.add_run(block.meta)
        x.italic = True
        x.font.size = Pt(style.contact_pt)
        x.font.name = style.font
    for b in block.bullets:
        bullet(doc, b, style)


def _render_project(doc: Document, block: Block, style: TemplateStyle) -> None:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(style.entry_space_before)
    h.paragraph_format.space_after = Pt(0)
    t = h.add_run(block.headline)
    t.bold = True
    t.font.size = Pt(style.entry_pt)
    t.font.name = style.font
    if block.headline_tech:
        g = h.add_run("  |  " + block.headline_tech)
        g.font.size = Pt(style.contact_pt)
        g.italic = True
        g.font.name = style.font
    if block.text:
        d = doc.add_paragraph()
        d.paragraph_format.space_after = Pt(2)
        dr = d.add_run(block.text)
        dr.font.size = Pt(style.body_pt)
        dr.font.name = style.font
    for det in block.bullets:
        bullet(doc, det, style)


BLOCK_RENDERERS = {
    "para": _render_para,
    "skillline": _render_skillline,
    "line": lambda doc, b, s: bullet(doc, b.text, s),
    "job": lambda doc, b, s: _render_job_or_edu(doc, b, s, "job"),
    "edu": lambda doc, b, s: _render_job_or_edu(doc, b, s, "edu"),
    "project": _render_project,
}


def _render_header(doc: Document, layout: Layout) -> None:
    """Name, professional title and contact line (text + hyperlinks)."""
    style = layout.style
    align = _align(style.header_align)
    accent = _accent(style)

    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(layout.name)
    r.bold = True
    r.font.size = Pt(style.name_pt)
    r.font.name = style.font
    p.paragraph_format.space_after = Pt(0)

    if layout.title:
        t = doc.add_paragraph()
        t.alignment = align
        r = t.add_run(layout.title)
        r.font.size = Pt(style.title_pt)
        r.font.color.rgb = accent
        r.font.name = style.font
        t.paragraph_format.space_after = Pt(4)

    line = doc.add_paragraph()
    line.alignment = align
    line.paragraph_format.space_after = Pt(2)

    def txt(s: str) -> None:
        x = line.add_run(s)
        x.font.name = style.font
        x.font.size = Pt(style.contact_pt)

    # Separators only BETWEEN items (layout already dropped empties).
    for i, (kind, text, url) in enumerate(layout.contacts):
        if i > 0:
            txt("  |  ")
        if kind == "link":
            add_hyperlink(line, text, url, style)
        else:
            txt(text)


def build_cv(
    cv: Dict[str, Any],
    template_key: str | None = None,
    profile_key: str | None = None,
) -> io.BytesIO:
    """Main entry: CV dict + template + profile -> BytesIO docx."""
    layout = build_layout(cv, template_key, profile_key)
    style = layout.style

    doc = Document()
    style_document(doc, style)

    # Word metadata (keywords come from actual user skills, capped at the
    # 255-char metadata limit Word enforces).
    core = doc.core_properties
    core.title = layout.name + " - CV"
    core.author = layout.name
    core.subject = layout.title
    skills_flat = ", ".join(
        str(item)
        for g in cv.get("skills", [])
        if isinstance(g, dict)
        for item in g.get("items", [])
    )
    core.keywords = skills_flat[:255] if skills_flat else layout.title

    _render_header(doc, layout)

    # ---- Body sections in PROFILE order (from the shared layout) ----
    for section in layout.sections:
        heading(doc, section.heading, style)
        for block in section.blocks:
            BLOCK_RENDERERS[block.kind](doc, block, style)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def generate_docx_bytes(
    cv: Dict[str, Any],
    template_key: str | None = None,
    profile_key: str | None = None,
) -> bytes:
    """CV dict + template + profile -> raw DOCX bytes."""
    return build_cv(cv, template_key, profile_key).getvalue()
